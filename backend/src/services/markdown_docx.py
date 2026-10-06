"""Markdown to a styled Word document, for work item documents.

Covers what the assistant and people write in these documents: headings,
paragraphs with **bold**, *italic* and `code`, bulleted and numbered
lists, fenced code blocks, tables and horizontal rules. Anything else is
kept as plain text rather than dropped.
"""

import io
import re

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")


def _runs(paragraph, text: str) -> None:
    from docx.shared import Pt

    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def markdown_to_docx(title: str, markdown: str, *, landscape: bool = False) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Cm(2))

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    lines = markdown.replace("\r\n", "\n").split("\n")
    if not any(line.startswith("# ") for line in lines[:3]):
        document.add_heading(title, level=0)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            paragraph = document.add_paragraph()
            run = paragraph.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and _TABLE_RULE.match(lines[i + 1]):
            header = _cells(line)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            table = document.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Light Grid Accent 1"
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                cell.paragraphs[0].add_run(text).bold = True
            for r, row in enumerate(rows, start=1):
                for c in range(len(header)):
                    cell = table.rows[r].cells[c]
                    cell.text = ""
                    _runs(cell.paragraphs[0], row[c] if c < len(row) else "")
            document.add_paragraph()
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            # The first "# " heading is the document's title; later ones are sections.
            level = len(heading.group(1))
            first_title = level == 1 and not any(earlier.startswith("# ") for earlier in lines[:i])
            document.add_heading(heading.group(2), level=0 if first_title else level)
        elif re.match(r"^[-*+]\s+", stripped):
            _runs(document.add_paragraph(style="List Bullet"), re.sub(r"^[-*+]\s+", "", stripped))
        elif re.match(r"^\d+[.)]\s+", stripped):
            _runs(document.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", stripped))
        elif re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            document.add_paragraph("_" * 60)
        elif stripped.startswith(">"):
            paragraph = document.add_paragraph()
            _runs(paragraph, stripped.lstrip("> "))
            paragraph.paragraph_format.left_indent = Cm(0.8)
        elif stripped:
            _runs(document.add_paragraph(), stripped)
        i += 1

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
