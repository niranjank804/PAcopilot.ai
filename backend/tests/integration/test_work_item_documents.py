"""Working a PBI without a file system: the assistant records the PBI,
writes its documents onto the work item, and links the changes it drafted;
people read, edit and download them."""

import io
import json

import pytest

from src.ai.tools.work_items import (
    GetWorkItemTool,
    LinkWorkItemChangeTool,
    SaveWorkItemDocumentTool,
    SaveWorkItemTool,
)
from src.services.markdown_docx import markdown_to_docx
from tests.integration.test_team_collaboration import team  # noqa: F401 - fixture
from tests.integration.tm1.test_changes_api import (  # noqa: F401 - fixtures
    fake_tm1_client,
    tm1_credentials_key,
)
from tests.integration.tm1.test_environments import _connection, _draft

PBI = (
    "Create a text-file based copy for SCN Workforce Planning Summary.\n\n"
    "Acceptance criteria:\n1. A new TI process is created.\n2. Period DIM handling is retained."
)


async def _tool(tool, db_session, team, **kwargs):
    db_session.info.pop("organization_id", None)
    return json.loads(await tool.execute(db_session, organization_id=team["org_id"], user_id=team["a_id"], **kwargs))


@pytest.mark.asyncio
async def test_the_assistant_records_a_pasted_pbi_and_reads_it_back(client, db_session, team):
    saved = await _tool(SaveWorkItemTool(), db_session, team,
                        reference="PBI 4076293", title="Workforce Planning Summary text copy", description=PBI)
    assert saved["created"] and saved["page"].startswith("/team/")

    again = await _tool(SaveWorkItemTool(), db_session, team,
                        reference="pbi 4076293", title="Workforce Planning Summary text copy",
                        description=PBI + "\n3. The Master process calls it.")
    assert again["created"] is False

    found = await _tool(GetWorkItemTool(), db_session, team, reference="PBI 4076293")
    assert found["found"] and "Master process" in found["description"]


@pytest.mark.asyncio
async def test_documents_are_versioned_listed_and_downloadable(client, db_session, team):
    await _tool(SaveWorkItemTool(), db_session, team, reference="PBI 1", title="T", description=PBI)
    first = await _tool(SaveWorkItemDocumentTool(), db_session, team, reference="PBI 1",
                        kind="clarification_email", title="Clarification email",
                        content="# Questions\n\n- Which **Period** elements?\n\n| Item | Default |\n|---|---|\n| A | B |")
    second = await _tool(SaveWorkItemDocumentTool(), db_session, team, reference="PBI 1",
                         kind="clarification_email", title="Clarification email", content="# Questions v2")
    assert (first["version"], second["version"]) == (1, 2)

    found = await _tool(GetWorkItemTool(), db_session, team, reference="PBI 1")
    assert found["documents"][0]["content"] == "# Questions v2"
    assert found["progress"]["Requirements clarified"] is True
    assert found["progress"]["Tested"] is False

    item_id = second["page"].split("/")[-1]
    listed = await client.get(f"/team/work-items/{item_id}/documents", headers=team["b"])
    assert listed.status_code == 200, listed.text
    [document] = listed.json()["data"]
    assert document["drafted_by_assistant"] is True

    word = await client.get(f"/team/work-items/{item_id}/documents/{document['id']}/download", headers=team["b"])
    assert word.status_code == 200
    assert word.content[:2] == b"PK"  # a .docx is a zip
    assert "PBI 1 - Clarification email.docx" in word.headers["content-disposition"]

    markdown = await client.get(
        f"/team/work-items/{item_id}/documents/{document['id']}/download?format=md", headers=team["b"]
    )
    assert markdown.text == "# Questions v2"


@pytest.mark.asyncio
async def test_people_edit_and_delete_documents(client, db_session, team):
    await _tool(SaveWorkItemTool(), db_session, team, reference="PBI 2", title="T")
    resp = await client.get("/team/work-items?q=PBI 2", headers=team["a"])
    item_id = next(i["id"] for i in resp.json()["data"] if i["reference"] == "PBI 2")

    saved = await client.post(f"/team/work-items/{item_id}/documents",
                              json={"kind": "test_results", "title": "Test results", "content": "All passed."},
                              headers=team["a"])
    assert saved.status_code == 201, saved.text
    assert saved.json()["data"]["drafted_by_assistant"] is False

    bad = await client.post(f"/team/work-items/{item_id}/documents",
                            json={"kind": "nonsense", "title": "x", "content": "y"}, headers=team["a"])
    assert bad.status_code == 422

    gone = await client.delete(f"/team/work-items/{item_id}/documents/{saved.json()['data']['id']}",
                               headers=team["b"])
    assert gone.status_code == 200
    assert (await client.get(f"/team/work-items/{item_id}/documents", headers=team["a"])).json()["data"] == []


@pytest.mark.asyncio
async def test_another_organization_cannot_read_the_documents(client, db_session, team):
    from tests.fixtures.factories import auth_headers, create_org_admin

    await _tool(SaveWorkItemTool(), db_session, team, reference="PBI 3", title="T")
    saved = await _tool(SaveWorkItemDocumentTool(), db_session, team, reference="PBI 3",
                        kind="notes", title="Notes", content="secret")
    item_id = saved["page"].split("/")[-1]
    _org, stranger = await create_org_admin(db_session)

    resp = await client.get(f"/team/work-items/{item_id}/documents", headers=auth_headers(stranger))
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_a_drafted_change_links_to_the_pbi(client, db_session, team, tm1_credentials_key, fake_tm1_client):
    await _tool(SaveWorkItemTool(), db_session, team, reference="PBI 4", title="T")
    connection = await _connection(client, team["a"], "dev")
    change = await _draft(client, team["a"], connection)

    linked = await _tool(LinkWorkItemChangeTool(), db_session, team, reference="PBI 4", change_id=change)
    assert linked["linked"]
    again = await _tool(LinkWorkItemChangeTool(), db_session, team, reference="PBI 4", change_id=change)
    assert again["already"]

    found = await _tool(GetWorkItemTool(), db_session, team, reference="PBI 4")
    assert found["progress"]["Proposed fix"] is True


def test_markdown_becomes_a_word_document_with_tables_lists_and_code():
    from docx import Document

    data = markdown_to_docx("Delivery", (
        "# Delivery document\n\n## Summary\n\nSome **bold** and `code`.\n\n"
        "- one\n- two\n\n1. first\n\n```\nCellPutN(1, 'Cube', 'A');\n```\n\n"
        "| Test | Result |\n|---|---|\n| Rerun | Identical |\n"
    ), landscape=True)
    document = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Delivery document" in text and "CellPutN" in text
    assert document.tables[0].cell(1, 1).text == "Identical"
    section = document.sections[0]
    assert section.page_width > section.page_height
